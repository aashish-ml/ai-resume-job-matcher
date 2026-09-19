from pathlib import Path

from docx import Document

from src.utils.logger import logger
from src.utils.validators import (
    FileValidationError,
    validate_uploaded_file,
)


class DOCXExtractionError(Exception):
    """Raised when DOCX text extraction fails."""


def extract_text_from_docx(file_path: str | Path) -> str:
    """
    Extract text from a DOCX resume.

    Parameters
    ----------
    file_path : str | Path
        Path to the DOCX document.

    Returns
    -------
    str
        Extracted text from paragraphs and tables.

    Raises
    ------
    FileNotFoundError
        If the DOCX file does not exist.

    FileValidationError
        If the file is invalid or unsupported.

    DOCXExtractionError
        If DOCX extraction fails.
    """

    path = Path(file_path)

    # ---------------------------------------------------------
    # File existence
    # ---------------------------------------------------------

    if not path.exists():
        raise FileNotFoundError(
            f"DOCX file not found: {path}"
        )

    if not path.is_file():
        raise FileValidationError(
            "Provided DOCX path is not a file."
        )

    # ---------------------------------------------------------
    # File validation
    # ---------------------------------------------------------

    file_size = path.stat().st_size

    extension = validate_uploaded_file(
        filename=path.name,
        file_size=file_size,
    )

    if extension != ".docx":
        raise FileValidationError(
            "DOCX parser can only process .docx files."
        )

    logger.info(
        "Starting DOCX text extraction: %s",
        path.name,
    )

    # ---------------------------------------------------------
    # DOCX extraction
    # ---------------------------------------------------------

    try:
        document = Document(str(path))

        extracted_parts: list[str] = []

        # -----------------------------------------------------
        # Paragraphs
        # -----------------------------------------------------

        for paragraph in document.paragraphs:
            text = paragraph.text.strip()

            if text:
                extracted_parts.append(text)

        # -----------------------------------------------------
        # Tables
        # -----------------------------------------------------

        for table in document.tables:
            for row in table.rows:
                row_text = []

                for cell in row.cells:
                    cell_text = cell.text.strip()

                    if cell_text:
                        row_text.append(cell_text)

                if row_text:
                    extracted_parts.append(
                        " | ".join(row_text)
                    )

        # -----------------------------------------------------
        # Combine extracted content
        # -----------------------------------------------------

        extracted_text = "\n".join(
            extracted_parts
        ).strip()

        if not extracted_text:
            raise DOCXExtractionError(
                "No readable text was extracted from the DOCX."
            )

        logger.info(
            "DOCX extraction completed: %s (%d characters)",
            path.name,
            len(extracted_text),
        )

        return extracted_text

    except DOCXExtractionError:
        raise

    except Exception as exc:
        logger.exception(
            "DOCX extraction failed for %s",
            path.name,
        )

        raise DOCXExtractionError(
            f"Failed to extract text from DOCX: {exc}"
        ) from exc